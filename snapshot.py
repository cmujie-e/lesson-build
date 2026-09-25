"""Text snapshot of a lesson's deliverables, for regression tests.

Usage:
  python snapshot.py save <lesson_folder> <snapshot.json>
  python snapshot.py compare <lesson_folder> <snapshot.json>

Records, for each deliverable, its extracted text (slides with notes, Word text, sheet
cells, PDF text) plus page/slide counts. `compare` reports any file whose text changed.
Binary differences (timestamps, zip ordering) are ignored on purpose: the test is whether
the content students and teachers see is the same.
"""
import json
import os
import re
import shutil
import subprocess
import sys

from toolpaths import PDFTOTEXT  # poppler, not Git's xpdf copy (see toolpaths.py)


def text_of(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        out = subprocess.run([PDFTOTEXT, "-enc", "UTF-8", path, "-"], capture_output=True, text=True, encoding="utf-8").stdout
        pages = out.count("\f")
    elif ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(path)
        parts = []
        for s in prs.slides:
            for sh in s.shapes:
                if sh.has_text_frame:
                    parts.append(sh.text_frame.text)
                elif getattr(sh, "has_table", False) and sh.has_table:
                    parts.append(" | ".join(c.text for r in sh.table.rows for c in r.cells))
                elif sh.shape_type == 13:
                    parts.append(f"[picture {sh.image.filename if hasattr(sh, 'image') else ''} {len(sh.image.blob)}]")
            if s.has_notes_slide:
                parts.append("NOTES: " + s.notes_slide.notes_text_frame.text)
        out, pages = "\n".join(parts), len(prs.slides)
    elif ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path)
        out = "\n".join(" | ".join("" if c is None else str(c) for c in r) for ws in wb for r in ws.iter_rows(values_only=True))
        pages = len(wb.sheetnames)
    else:
        from markitdown import MarkItDown
        out, pages = MarkItDown().convert(path).text_content, None
    return re.sub(r"[ \t]+", " ", out).strip(), pages


def snapshot(folder):
    files = sorted(f for f in os.listdir(folder) if os.path.isfile(os.path.join(folder, f))
                   and not f.startswith("~$") and f != "content.md")
    return {f: dict(zip(("text", "pages"), text_of(os.path.join(folder, f)))) for f in files}


if __name__ == "__main__":
    mode, folder, snap = sys.argv[1:4]
    if mode == "save":
        json.dump(snapshot(folder), open(snap, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        print(f"saved {snap}")
    else:
        old, new = json.load(open(snap, encoding="utf-8")), snapshot(folder)
        changed = False
        for f in sorted(set(old) | set(new)):
            if f not in new or f not in old:
                print(f"[DIFF] {f}: {'missing now' if f not in new else 'new file'}")
                changed = True
            elif old[f] != new[f]:
                a, b = old[f]["text"], new[f]["text"]
                i = next((k for k in range(min(len(a), len(b))) if a[k] != b[k]), min(len(a), len(b)))
                print(f"[DIFF] {f}: pages {old[f]['pages']} -> {new[f]['pages']}; first change at char {i}:")
                print(f"       was: ...{a[max(0, i - 60):i + 80]!r}")
                print(f"       now: ...{b[max(0, i - 60):i + 80]!r}")
                changed = True
            else:
                print(f"[SAME] {f}")
        print("\nRESULT:", "CHANGED" if changed else "IDENTICAL")
        sys.exit(1 if changed else 0)
