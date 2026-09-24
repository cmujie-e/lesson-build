"""Build the bilingual glossary workbook from lesson.json.

Usage: python build_glossary.py <lesson.json> <out.xlsx>
Sheet 1 "Glossary": navy header, frozen header row, alternating shading, landscape, fit to width.
Sheet 2 "Notes": translation-accuracy notice and consistency notes.
"""
import json
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.properties import PageSetupProperties

NAVY, LIGHT = "1E2761", "F4F6FC"


def main(src, dst):
    with open(src, encoding="utf-8") as f:
        L = json.load(f)
    rows = L["glossary"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Glossary"
    for r in rows:
        ws.append(r)
    widths = [30, 62, 34, 30]
    for i, w in enumerate(widths):
        ws.column_dimensions["ABCD"[i]].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF", name="Calibri", size=12)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for ri, row in enumerate(ws.iter_rows(min_row=2), start=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.font = Font(name="Calibri", size=11, bold=(c.column == 1))
            if ri % 2 == 1:
                c.fill = PatternFill("solid", fgColor=LIGHT)
    ws.freeze_panes = "A2"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = "1:1"

    notes = wb.create_sheet("Notes")
    notes.column_dimensions["A"].width = 100
    notes.page_setup.orientation = "landscape"
    notes.page_setup.fitToWidth, notes.page_setup.fitToHeight = 1, 0
    notes.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    for i, line in enumerate(L["glossary_notes"], start=1):
        c = notes.cell(row=i, column=1, value=line)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(name="Calibri", size=12 if i == 1 else 11, bold=(i == 1), color=NAVY if i == 1 else "000000")
    wb.save(dst)
    print(f"  {dst.split(chr(92))[-1]} ({len(rows) - 1} terms)")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
