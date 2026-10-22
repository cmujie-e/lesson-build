"""Build a teacher guide (.docx) from a light Markdown source, in the lesson-build house style.

Usage: python build_guide.py <guide.md> <out.docx>

Source format (one block per line):
  # Title                      cover title (first line); "subtitle: ..." on the next line
  ## Part heading              navy banner; "Part 2" and "Part 3" begin a new page
  ### Sub-heading              navy bold heading
  #### Slide N · Title [TAG]   slide bar (lighter navy)
  - text                       bullet (two-space indent "  - " = second level)
  1. text                      numbered line (number kept as typed)
  | a | b |                    table; first row is the header
  !LABEL: text                 callout box; LABEL one of SAY ASK EXPECT WATCH WRONG FURTHER EXAM FIX CALC NOTE
                               " // " inside a callout starts a new line
  ```  ...  ```                monospace block (worked calculations, truth tables)
  [[pagebreak]]
  plain line                   paragraph; **bold** and *italic* inline
"""
import re
import sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

NAVY, DOT, INK = "1E2761", "3A4488", RGBColor(0x1E, 0x27, 0x61)
CALLOUTS = {
    "SAY": ("Say", "EEF3FF"), "ASK": ("Ask", "FFF6E0"), "EXPECT": ("Expect", "EAF7EE"),
    "WATCH": ("Watch out", "FFEAD1"), "WRONG": ("If they get it wrong", "FFF0EC"),
    "FURTHER": ("Going further (beyond AS)", "F4F6FC"), "EXAM": ("Exam technique", "FFF6E0"),
    "FIX": ("Textbook correction", "FFF0EC"), "CALC": ("Worked calculation", "F4F6FC"),
    "NOTE": ("Note", "F4F6FC"), "TIME": ("Timing", "F4F6FC"),
}


def shade(cell_or_par, fill):
    pr = cell_or_par._tc.get_or_add_tcPr() if hasattr(cell_or_par, "_tc") else cell_or_par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    pr.append(shd)


def runs(par, text, size=10.5, color=None, bold=False, font=None):
    """Add text with **bold** and *italic* markup."""
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        b, i = bold, False
        if part.startswith("**") and part.endswith("**"):
            part, b = part[2:-2], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, i = part[1:-1], True
        r = par.add_run(part)
        r.bold, r.italic = b, i
        r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
        if font:
            r.font.name = font
            r.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    return par


def spacing(par, before=0, after=4, keep=False):
    pf = par.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if keep:
        pf.keep_with_next = True


def bar(doc, text, fill, size, before=10):
    p = doc.add_paragraph()
    shade(p, fill)
    spacing(p, before, 6, keep=True)
    runs(p, " " + text, size=size, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)


def callout(doc, label, text):
    name, fill = CALLOUTS[label]
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    shade(c, fill)
    trPr = t.rows[0]._tr.get_or_add_trPr()  # keep each callout box whole on one page
    cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)
    lines = text.split(" // ")
    p = c.paragraphs[0]
    spacing(p, 2, 2)
    runs(p, f"{name}: ", size=10, color=INK, bold=True)
    runs(p, lines[0], size=10)
    for ln in lines[1:]:
        q = c.add_paragraph()
        spacing(q, 0, 2)
        runs(q, ln, size=10)
    gap = doc.add_paragraph()
    spacing(gap, 0, 2)


def table(doc, rows):
    cols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=cols)
    t.style = "Table Grid"
    for row in t.rows:  # never split a row across pages
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true")
        row._tr.get_or_add_trPr().append(cs)
    for i, r in enumerate(rows):
        for j in range(cols):
            c = t.cell(i, j)
            p = c.paragraphs[0]
            spacing(p, 1, 1)
            txt = r[j] if j < len(r) else ""
            if i == 0:
                if j == 0:
                    trPr = t.rows[0]._tr.get_or_add_trPr()
                    h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
                shade(c, DOT)
                runs(p, txt, size=9.5, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
            else:
                if i % 2 == 0:
                    shade(c, "F4F6FC")
                runs(p, txt, size=9.5)
    gap = doc.add_paragraph()
    spacing(gap, 0, 4)


def mono(doc, lines):
    t = doc.add_table(rows=1, cols=1)
    c = t.cell(0, 0)
    shade(c, "F4F6FC")
    first = True
    for ln in lines:
        p = c.paragraphs[0] if first else c.add_paragraph()
        first = False
        spacing(p, 0, 0)
        runs(p, ln if ln else " ", size=9.5, font="Courier New")
    gap = doc.add_paragraph()
    spacing(gap, 0, 4)


def footer(doc, text):
    sec = doc.sections[0]
    p = sec.footer.paragraphs[0]
    runs(p, text + "  ·  page ", size=8, color=INK)
    r = p.add_run()
    for tag, val in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = val
        r._r.append(el)
    r.font.size = Pt(8)


def set_compat_mode(doc, level="15"):
    """python-docx's template has no compatibilityMode setting, so Word opens the file in
    'Compatibility Mode'. Declare Word 2013+ (15) so the title bar shows no such label."""
    settings = doc.settings.element
    compat = settings.find(qn("w:compat"))
    if compat is None:
        compat = OxmlElement("w:compat")
        settings.append(compat)
    for s in compat.findall(qn("w:compatSetting")):
        if s.get(qn("w:name")) == "compatibilityMode":
            compat.remove(s)
    cs = OxmlElement("w:compatSetting")
    cs.set(qn("w:name"), "compatibilityMode")
    cs.set(qn("w:uri"), "http://schemas.microsoft.com/office/word")
    cs.set(qn("w:val"), level)
    compat.append(cs)


def build(src, out):
    lines = open(src, encoding="utf-8").read().splitlines()
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.font.size = Pt(10.5)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.9)
    sec.top_margin = sec.bottom_margin = Cm(1.7)
    i, first_part = 0, True
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            i += 1
            continue
        if ln.startswith("# "):
            bar(doc, ln[2:], NAVY, 18, before=0)
            if i + 1 < len(lines) and lines[i + 1].startswith("subtitle:"):
                p = doc.add_paragraph(); spacing(p, 2, 10)
                runs(p, lines[i + 1][9:].strip(), size=11, color=INK)
                footer(doc, lines[i + 1][9:].strip().split("·")[0].strip() + " · " + ln[2:])
                i += 1
        elif ln.startswith("#### "):
            bar(doc, ln[5:], DOT, 11.5, before=10)
        elif ln.startswith("### "):
            p = doc.add_paragraph(); spacing(p, 8, 3, keep=True)
            runs(p, ln[4:], size=12, color=INK, bold=True)
        elif ln.startswith("## "):
            # new page only for the two parts that are printed separately: subject knowledge and the walkthrough
            if re.match(r"Part [23] ", ln[3:]) and not first_part:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            first_part = False
            bar(doc, ln[3:], NAVY, 14, before=6)
        elif ln.startswith("[[pagebreak]]"):
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif ln.startswith("```"):
            block = []
            i += 1
            while not lines[i].startswith("```"):
                block.append(lines[i].rstrip())
                i += 1
            mono(doc, block)
        elif ln.lstrip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            table(doc, rows)
            continue
        elif re.match(r"^!([A-Z]+):\s", ln):
            m = re.match(r"^!([A-Z]+):\s(.*)$", ln)
            callout(doc, m.group(1), m.group(2))
        elif re.match(r"^\s*- ", ln):
            level = 1 if ln.startswith("  ") else 0
            p = doc.add_paragraph()
            spacing(p, 0, 2)
            p.paragraph_format.left_indent = Cm(0.6 + 0.6 * level)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            runs(p, ("•  " if level == 0 else "–  ") + ln.strip()[2:])
        elif re.match(r"^\d+\.\s", ln):
            p = doc.add_paragraph()
            spacing(p, 0, 2)
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            runs(p, ln)
        else:
            p = doc.add_paragraph()
            spacing(p, 0, 5)
            runs(p, ln)
        i += 1
    set_compat_mode(doc)
    doc.save(out)
    print("saved", out)


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
