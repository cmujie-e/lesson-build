"""Per-slide check against Chapter 3 CLAUDE.md.

Reports, for every slide: content word count (limit 35, title excluded), whether speaker
notes exist, whether question slides' notes contain an answer, and content font sizes (rule: 30 pt).

Usage: python check_deck.py <deck.pptx> [<deck.pptx> ...]

What counts as content: every text shape and table except the slide chrome.
Chrome = shapes named Title / Tag / Footer / PageNum (set by build_deck.js). For decks built
elsewhere without those names, falls back to position: title bar zone (top < 1.2 in) and
footer zone (top >= 7.0 in), plus shapes whose whole text is a bare number.
A "word" is a whitespace-separated token containing at least one letter or digit, so
separators like "•" or "–" are not counted.
Question slides are detected by their tag (DO NOW, CFU, STOP & CHECK, STRETCH, EXIT TICKET);
their notes must contain "Answer:".
Every slide with a picture must have "Image:" in its notes and a "Slide N:" line on the
closing credits slide (tag CREDITS), which is itself exempt from the word and font rules.
"""
import re
import sys
from pptx import Presentation
from pptx.util import Inches

WORD_LIMIT = 35
FONT_PT = 30
TITLE_ZONE = Inches(1.2)
FOOTER_ZONE = Inches(7.0)
WORD = re.compile(r"[A-Za-z0-9]")
CHROME = {"Title", "Tag", "Footer", "PageNum"}
QUESTION_TAGS = {"DO NOW", "CFU", "STOP & CHECK", "STRETCH", "EXIT TICKET"}


def text_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # group
            yield from text_shapes(sh.shapes)
        elif getattr(sh, "has_table", False) and sh.has_table:
            yield sh
        elif sh.has_text_frame and sh.text_frame.text.strip():
            yield sh


def frames(sh):
    if getattr(sh, "has_table", False) and sh.has_table:
        return [c.text_frame for r in sh.table.rows for c in r.cells]
    return [sh.text_frame]


def is_content(sh, named):
    if named:
        return sh.name not in CHROME
    if sh.top is not None and (sh.top < TITLE_ZONE or sh.top >= FOOTER_ZONE):
        return False
    return not " ".join(f.text for f in frames(sh)).strip().isdigit()


def pictures(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from pictures(sh.shapes)
        elif sh.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
            yield sh


def check(path):
    prs = Presentation(path)
    print(f"\n{path}")
    print(f"{'Slide':>5} | {'Words':>5} | {'<=35':<4} | {'Notes':<5} | {'Answer':<6} | {'Image':<6} | {'Content pt':<14} | Title")
    print("-" * 113)
    over, missing, no_answer, off_font, uncredited = [], [], [], [], []
    credits_text = ""
    for slide in prs.slides:
        if any(s.name == "Tag" and s.text_frame.text.strip() == "CREDITS" for s in text_shapes(slide.shapes)):
            credits_text = " ".join(f.text for s in text_shapes(slide.shapes) for f in frames(s))
    for i, slide in enumerate(prs.slides, 1):
        shapes = list(text_shapes(slide.shapes))
        named = any(s.name in CHROME for s in shapes)
        # the closing credits slide is exempt from the word and font rules (CLAUDE.md)
        exempt = any(s.name == "Tag" and s.text_frame.text.strip() == "CREDITS" for s in shapes)
        content = [] if exempt else [s for s in shapes if is_content(s, named)]
        words = sum(len([w for w in f.text.split() if WORD.search(w)]) for s in content for f in frames(s))
        sizes = sorted({r.font.size.pt for s in content for f in frames(s)
                        for p in f.paragraphs for r in p.runs if r.font.size and r.text.strip()})
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        title = next((s.text_frame.text for s in shapes if s.name == "Title"), None)
        if title is None:
            title = next((s.text_frame.text for s in shapes if s.top is not None and s.top < TITLE_ZONE and s.has_text_frame), "")
        tag = next((s.text_frame.text.strip() for s in shapes if s.name == "Tag"), "")
        is_q = tag in QUESTION_TAGS
        has_answer = "Answer:" in notes
        if words > WORD_LIMIT:
            over.append(i)
        if not notes:
            missing.append(i)
        if is_q and not has_answer:
            no_answer.append(i)
        if any(sz != FONT_PT for sz in sizes):
            off_font.append(i)
        n_pics = len(list(pictures(slide.shapes)))
        credited = "Image:" in notes and f"Slide {i}:" in credits_text
        if n_pics and not credited:
            uncredited.append(i)
        ans = ("yes" if has_answer else "NO") if is_q else "-"
        img = ("yes" if credited else "NO") if n_pics else "-"
        size_txt = "exempt" if exempt else (", ".join(f"{sz:g}" for sz in sizes) or "-")
        print(f"{i:>5} | {'-' if exempt else words:>5} | {'-' if exempt else ('OK' if words <= WORD_LIMIT else 'OVER'):<4} | "
              f"{'yes' if notes else 'NO':<5} | {ans:<6} | {img:<6} | {size_txt:<14} | {title.replace(chr(10), ' ')[:42]}")
    print("-" * 113)
    print(f"{len(prs.slides)} slides")
    print(f"  Over {WORD_LIMIT} words (title excluded) : {len(over)}  {over}")
    print(f"  Missing speaker notes         : {len(missing)}  {missing}")
    print(f"  Question slides with no answer: {len(no_answer)}  {no_answer}")
    print(f"  Content text not {FONT_PT} pt        : {len(off_font)}  {off_font}")
    print(f"  Images without credit         : {len(uncredited)}  {uncredited}  (Image column: credit in notes AND on credits slide)")
    return not (over or missing or no_answer or off_font or uncredited)


if __name__ == "__main__":
    ok = all([check(p) for p in sys.argv[1:]])
    print("\nRESULT:", "PASS" if ok else "FAIL")
    sys.exit(0 if ok else 1)
